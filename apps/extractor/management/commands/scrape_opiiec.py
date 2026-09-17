"""
Management command : scraping du référentiel OPIIEC (Observatoire des métiers
du Numérique, de l'Ingénierie, du Conseil et de l'Événement).

Sources :
  - Fiches métiers (168 profils) via export CSV
  - Référentiel de compétences (87 macro-compétences, 348 compétences détaillées)
  - Certifications OPIIEC (408 certifications)

Données publiques : https://www.opiiec.fr/metiers
"""
import csv
import io
import logging
import re
import time
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from django.core.management.base import BaseCommand
from django.utils import timezone

from apps.extractor.models import JobPosting, FormationCatalog
from apps.extractor.classifier import (
    classify_training_type,
    classify_sector,
    compute_relevance_score,
    extract_skills,
)

logger = logging.getLogger(__name__)

BASE_URL = 'https://www.opiiec.fr'

HEADERS = {
    'User-Agent': (
        'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) '
        'AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36'
    ),
    'Accept-Language': 'fr-FR,fr;q=0.9,en;q=0.8',
}

# Known OPIIEC fiche IDs (discovered from public exports)
# Range is approximately 82940-83120 with gaps
KNOWN_FICHE_IDS = list(range(82940, 83120))

# AI/Data/Digital keywords to filter relevant fiches
AI_RELEVANCE_KEYWORDS = [
    'intelligence artificielle', 'ia', 'machine learning', 'deep learning',
    'data', 'donnée', 'données', 'numérique', 'digital', 'cloud',
    'cybersécurité', 'cyber', 'sécurité informatique',
    'développeur', 'développement', 'logiciel', 'devops',
    'infrastructure', 'réseau', 'système', 'iot', 'objets connectés',
    'blockchain', 'big data', 'analytics', 'bi ', 'business intelligence',
    'erp', 'crm', 'saas', 'agile', 'scrum', 'product owner',
    'chef de projet', 'architecte', 'consultant', 'ingénieur',
    'automatisation', 'robotique', 'rpa',
]


class Command(BaseCommand):
    help = 'Scrape le référentiel OPIIEC : fiches métiers + compétences'

    def add_arguments(self, parser):
        parser.add_argument(
            '--source',
            type=str,
            default='all',
            choices=['all', 'metiers', 'competences', 'certifications'],
            help='Source à scraper (default: all)',
        )
        parser.add_argument(
            '--limit',
            type=int,
            default=0,
            help='Nombre max de fiches (0 = toutes)',
        )

    def handle(self, *args, **options):
        source = options['source']
        limit = options['limit']
        total = 0

        if source in ('all', 'metiers'):
            self.stdout.write(self.style.NOTICE('[OPIIEC] Scraping des fiches métiers...'))
            count = self._scrape_metiers(limit)
            total += count
            self.stdout.write(self.style.SUCCESS(f'  → {count} fiches métiers importées'))

        if source in ('all', 'competences'):
            self.stdout.write(self.style.NOTICE('[OPIIEC] Scraping du référentiel de compétences...'))
            count = self._scrape_competences()
            total += count
            self.stdout.write(self.style.SUCCESS(f'  → {count} compétences extraites'))

        if source in ('all', 'certifications'):
            self.stdout.write(self.style.NOTICE('[OPIIEC] Scraping des certifications...'))
            count = self._scrape_certifications(limit)
            total += count
            self.stdout.write(self.style.SUCCESS(f'  → {count} certifications importées'))

        self.stdout.write(self.style.SUCCESS(f'\n[OPIIEC] Total : {total} éléments importés'))

    # ─── Fiches métiers ───────────────────────────────────────────────

    def _scrape_metiers(self, limit):
        """Scrape toutes les fiches métiers OPIIEC via export CSV."""
        # Step 1: Discover fiche IDs from listing pages
        fiche_ids = self._discover_fiche_ids()
        if not fiche_ids:
            self.stdout.write(self.style.WARNING(
                '  Impossible de découvrir les IDs, utilisation des IDs connus...'
            ))
            fiche_ids = KNOWN_FICHE_IDS

        if limit > 0:
            fiche_ids = fiche_ids[:limit]

        self.stdout.write(f'  {len(fiche_ids)} fiches à traiter...')
        count = 0
        errors = 0

        for i, fiche_id in enumerate(fiche_ids):
            try:
                result = self._fetch_fiche_csv(fiche_id)
                if result:
                    count += result
                if (i + 1) % 20 == 0:
                    self.stdout.write(f'  ... {i + 1}/{len(fiche_ids)} traitées ({count} importées)')
                time.sleep(0.3)  # Rate limiting
            except Exception as e:
                errors += 1
                if errors <= 5:
                    logger.warning(f'OPIIEC fiche {fiche_id}: {e}')

        return count

    def _discover_fiche_ids(self):
        """Découvre tous les IDs de fiches depuis les pages de listing."""
        fiche_ids = []
        try:
            for page in range(18):  # ~17 pages of 10 items
                url = f'{BASE_URL}/metiers?page={page}'
                resp = requests.get(url, headers=HEADERS, timeout=15)
                if resp.status_code != 200:
                    break

                soup = BeautifulSoup(resp.text, 'html.parser')
                links = soup.find_all('a', href=re.compile(r'/metiers/(\d+)'))
                if not links:
                    break

                for link in links:
                    match = re.search(r'/metiers/(\d+)', link.get('href', ''))
                    if match:
                        fid = int(match.group(1))
                        if fid not in fiche_ids:
                            fiche_ids.append(fid)

                time.sleep(0.5)

        except Exception as e:
            logger.warning(f'OPIIEC discovery: {e}')

        self.stdout.write(f'  {len(fiche_ids)} fiches découvertes via le listing')
        return fiche_ids

    def _fetch_fiche_csv(self, fiche_id):
        """Télécharge et parse une fiche métier au format CSV."""
        url = f'{BASE_URL}/export_csv/{fiche_id}'
        try:
            resp = requests.get(url, headers=HEADERS, timeout=15)
            if resp.status_code != 200 or len(resp.content) < 50:
                return 0

            # OPIIEC CSV uses semicolon delimiter
            content = resp.content.decode('utf-8-sig', errors='replace')
            reader = csv.reader(io.StringIO(content), delimiter=';')
            rows = list(reader)

            if len(rows) < 2:
                return 0

            return self._parse_fiche_rows(rows, fiche_id)

        except requests.RequestException:
            return 0

    def _parse_fiche_rows(self, rows, fiche_id):
        """Parse les lignes CSV d'une fiche métier OPIIEC."""
        # CSV structure: headers in first row, data in second
        # Fields: Code, Referentiel, Date modification, Metier, Metier feminin,
        #         Famille, Appellations alternatives, Appellations anglaises,
        #         Finalite, Missions principales, Contexte, Conditions,
        #         Competences..., Certifications, Codes ROME, Mobilite
        headers = rows[0] if rows else []
        data = rows[1] if len(rows) > 1 else []

        if not headers or not data:
            return 0

        # Build dict from headers and data
        fiche = {}
        for i, header in enumerate(headers):
            if i < len(data):
                fiche[header.strip()] = data[i].strip()

        # Extract key fields
        title = (
            fiche.get('Metier', '')
            or fiche.get('Métier', '')
            or fiche.get('metier', '')
        )
        if not title:
            # Try to find title-like field
            for key, val in fiche.items():
                if 'metier' in key.lower() or 'métier' in key.lower():
                    if val and len(val) > 3:
                        title = val
                        break

        if not title:
            return 0

        family = fiche.get('Famille', '')
        finalite = fiche.get('Finalite', fiche.get('Finalité', ''))
        missions = fiche.get('Missions principales', '')
        contexte = fiche.get('Contexte organisationnel', '')
        appellations = fiche.get('Appellations alternatives', '')
        appellations_en = fiche.get('Appellations anglaises', '')
        certifications = fiche.get('Certifications', '')
        codes_rome = fiche.get('Codes ROME', '')

        # Collect all competences from the CSV
        competences = []
        for key, val in fiche.items():
            if any(kw in key.lower() for kw in ['compétence', 'competence', 'savoir']):
                if val and len(val) > 2:
                    competences.append(val)

        # Build rich description
        description_parts = []
        if family:
            description_parts.append(f'Famille : {family}')
        if finalite:
            description_parts.append(f'Finalité : {finalite}')
        if missions:
            description_parts.append(f'Missions : {missions}')
        if contexte:
            description_parts.append(f'Contexte : {contexte}')
        if appellations:
            description_parts.append(f'Appellations : {appellations}')
        if appellations_en:
            description_parts.append(f'English : {appellations_en}')
        if competences:
            description_parts.append(f'Compétences : {" | ".join(competences[:15])}')
        if certifications:
            description_parts.append(f'Certifications : {certifications}')
        if codes_rome:
            description_parts.append(f'Codes ROME : {codes_rome}')

        description = '\n'.join(description_parts)

        # Save as JobPosting (enriches skill demand data)
        return self._save_job(
            title=title,
            description=description,
            fiche_id=fiche_id,
            competences=competences,
        )

    def _save_job(self, title, description, fiche_id, competences=None):
        """Sauvegarde une fiche métier OPIIEC comme JobPosting."""
        full_text = f'{title} {description}'

        # Extract skills from both classifier and OPIIEC competences
        skills = extract_skills(full_text)

        # Also map OPIIEC competences to our skill taxonomy
        if competences:
            comp_text = ' '.join(competences)
            extra_skills = extract_skills(comp_text)
            skills = list(set(skills) | set(extra_skills))

        url = f'{BASE_URL}/metiers/{fiche_id}'

        _, created = JobPosting.objects.update_or_create(
            url=url,
            defaults={
                'title': title[:500],
                'company': 'OPIIEC (Référentiel)',
                'source': 'opiiec',
                'sector': classify_sector(full_text),
                'skills': skills,
                'location': 'France',
                'training_type_match': classify_training_type(full_text),
                'raw_description': description[:5000],
                'relevance_score': compute_relevance_score(full_text),
            },
        )
        return 1 if created else 0

    # ─── Référentiel de compétences ───────────────────────────────────

    def _scrape_competences(self):
        """Télécharge le référentiel de compétences OPIIEC (Excel)."""
        # Try the direct Excel download
        excel_url = (
            f'{BASE_URL}/sites/default/files/inline-files/'
            'Arbre-de-competences-OPIIEC-2024.xlsx'
        )

        try:
            resp = requests.get(excel_url, headers=HEADERS, timeout=30)
            if resp.status_code == 200 and len(resp.content) > 1000:
                return self._parse_competences_excel(resp.content)
        except Exception as e:
            logger.warning(f'OPIIEC competences Excel download: {e}')

        # Fallback: scrape the methodology page for competences
        self.stdout.write('  Excel non disponible, scraping de la page méthodologie...')
        return self._scrape_competences_page()

    def _parse_competences_excel(self, content):
        """Parse le fichier Excel des compétences OPIIEC."""
        try:
            import openpyxl
        except ImportError:
            self.stdout.write(self.style.WARNING(
                '  openpyxl non installé. Installez avec: pip install openpyxl'
            ))
            return self._scrape_competences_page()

        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
            count = 0

            for sheet in wb.sheetnames:
                ws = wb[sheet]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    if not row or not row[0]:
                        continue

                    # Extract competence info
                    macro_comp = str(row[0]).strip() if row[0] else ''
                    if not macro_comp or len(macro_comp) < 3:
                        continue

                    # Build a description from all columns
                    details = []
                    for cell in row[1:]:
                        if cell and str(cell).strip():
                            details.append(str(cell).strip())

                    description = ' | '.join(details) if details else ''
                    full_text = f'{macro_comp} {description}'

                    # Create a FormationCatalog entry for each competence area
                    # This enriches the "supply" side of gap analysis
                    _, created = FormationCatalog.objects.update_or_create(
                        name=f'Compétence OPIIEC : {macro_comp[:450]}',
                        provider='OPIIEC (Référentiel)',
                        defaults={
                            'description': description[:2000],
                            'sector': classify_sector(full_text),
                            'training_type': '',
                            'url': f'{BASE_URL}/methodologie',
                            'skills_covered': extract_skills(full_text),
                        },
                    )
                    if created:
                        count += 1

            wb.close()
            return count

        except Exception as e:
            logger.error(f'OPIIEC Excel parse error: {e}')
            return self._scrape_competences_page()

    def _scrape_competences_page(self):
        """Scrape les compétences depuis les pages web OPIIEC."""
        count = 0
        try:
            url = f'{BASE_URL}/recherche_competence'
            resp = requests.get(url, headers=HEADERS, timeout=15)
            if resp.status_code != 200:
                return 0

            soup = BeautifulSoup(resp.text, 'html.parser')

            # Extract competence blocks from the page
            comp_items = soup.find_all(['div', 'li', 'span'], class_=re.compile(
                r'competence|skill|comp-item|taxonomy', re.IGNORECASE
            ))

            for item in comp_items:
                text = item.get_text(strip=True)
                if text and len(text) > 5 and len(text) < 300:
                    _, created = FormationCatalog.objects.update_or_create(
                        name=f'Compétence OPIIEC : {text[:450]}',
                        provider='OPIIEC (Référentiel)',
                        defaults={
                            'description': f'Compétence identifiée dans le référentiel OPIIEC',
                            'url': url,
                            'skills_covered': extract_skills(text),
                        },
                    )
                    if created:
                        count += 1

        except Exception as e:
            logger.warning(f'OPIIEC competences page: {e}')

        return count

    # ─── Certifications ───────────────────────────────────────────────

    def _scrape_certifications(self, limit):
        """Scrape les certifications OPIIEC → FormationCatalog."""
        count = 0
        page = 0
        max_pages = 42  # ~408 certifs / 10 per page

        while page < max_pages:
            if limit > 0 and count >= limit:
                break

            try:
                url = f'{BASE_URL}/certifications?page={page}'
                resp = requests.get(url, headers=HEADERS, timeout=15)
                if resp.status_code != 200:
                    break

                soup = BeautifulSoup(resp.text, 'html.parser')

                # Find certification links
                cert_links = soup.find_all('a', href=re.compile(r'/certifications/'))
                if not cert_links:
                    break

                for link in cert_links:
                    href = link.get('href', '')
                    if not href or href == '/certifications':
                        continue

                    cert_url = urljoin(BASE_URL, href)
                    cert_name = link.get_text(strip=True)

                    if not cert_name or len(cert_name) < 5:
                        continue

                    # Extract additional info from parent elements
                    parent = link.find_parent(['div', 'li', 'tr'])
                    extra_text = parent.get_text(' ', strip=True) if parent else ''

                    # Determine certification type/level
                    training_type = 'certification'
                    if any(kw in extra_text.lower() for kw in ['master', 'diplôme', 'bac+']):
                        training_type = 'diplome'
                    elif any(kw in extra_text.lower() for kw in ['mooc', 'en ligne']):
                        training_type = 'mooc'

                    full_text = f'{cert_name} {extra_text}'

                    _, created = FormationCatalog.objects.update_or_create(
                        name=cert_name[:500],
                        provider='OPIIEC / France Compétences',
                        defaults={
                            'description': extra_text[:2000],
                            'sector': classify_sector(full_text),
                            'training_type': training_type,
                            'url': cert_url,
                            'skills_covered': extract_skills(full_text),
                        },
                    )
                    if created:
                        count += 1

                    if limit > 0 and count >= limit:
                        break

                page += 1
                time.sleep(0.5)

            except Exception as e:
                logger.warning(f'OPIIEC certifications page {page}: {e}')
                page += 1

        return count
